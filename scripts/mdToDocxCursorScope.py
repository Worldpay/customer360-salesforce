#!/usr/bin/env python3
"""Convert docs/cursorDevelopmentScope.md to .docx for Confluence paste."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.oxml import OxmlElement

REPO = Path(__file__).resolve().parents[1]
MD_PATH = REPO / "docs" / "cursorDevelopmentScope.md"
OUT_PATH = REPO / "docs" / "cursorDevelopmentScope.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rPr.insert(0, rFonts)


def add_formatted_paragraph(doc, text, style=None, mono=False):
    p = doc.add_paragraph(style=style)
    if not text:
        return p
    # Split on **bold** and `code` and [link](url)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        run = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            set_run_font(run, bold=True, size=11 if not mono else 10)
            run.text = part[2:-2]
        elif part.startswith("`") and part.endswith("`"):
            set_run_font(run, name="Consolas", size=10)
            run.text = part[1:-1]
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if m:
                set_run_font(run, size=11 if not mono else 10)
                run.text = f"{m.group(1)} ({m.group(2)})"
            else:
                set_run_font(run, size=11 if not mono else 10)
                run.text = part
        else:
            set_run_font(run, name="Consolas" if mono else "Calibri", size=10 if mono else 11)
            run.text = part
    if mono:
        for run in p.runs:
            run.font.name = "Consolas"
    return p


def parse_table_row(line):
    line = line.strip()
    if not line.startswith("|"):
        return None
    cells = [c.strip() for c in line.strip("|").split("|")]
    return cells


def is_separator_row(cells):
    if not cells:
        return False
    return all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells)


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            # Preserve "- " bullets as lines in cell
            lines = cell_text.split(" - ")
            if len(lines) > 1 and lines[0].strip() == "":
                lines = ["- " + x.strip() for x in cell_text.split("\n")]
            elif " - " in cell_text and cell_text.strip().startswith("-"):
                lines = [x.strip() for x in re.split(r"\s+-\s+", cell_text) if x.strip()]
                lines = [("- " + ln) if not ln.startswith("-") else ln for ln in lines]
            else:
                lines = cell_text.split("\n") if "\n" in cell_text else [cell_text]
            for li, ln in enumerate(lines):
                if li > 0:
                    p = cell.add_paragraph()
                add_formatted_paragraph_to_existing(p, ln)


def add_formatted_paragraph_to_existing(p, text):
    for run in list(p.runs):
        p._p.remove(run._r)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        run = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            set_run_font(run, bold=True)
            run.text = part[2:-2]
        elif part.startswith("`") and part.endswith("`"):
            set_run_font(run, name="Consolas", size=10)
            run.text = part[1:-1]
        else:
            set_run_font(run)
            run.text = part


def heading_level(line):
    m = re.match(r"^(#{1,6})\s+(.+)$", line)
    if not m:
        return None
    return len(m.group(1)), m.group(2).strip()


def convert(md_text: str, out_path: Path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lang = ""
    code_lines = []
    table_buffer = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            rows = []
            for tl in table_buffer:
                cells = parse_table_row(tl)
                if cells and not is_separator_row(cells):
                    rows.append(cells)
            add_table(doc, rows)
            table_buffer = []

    while i < len(lines):
        line = lines[i]
        raw = line

        if in_code:
            if line.strip().startswith("```"):
                in_code = False
                p = doc.add_paragraph()
                p.style = "No Spacing"
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, name="Consolas", size=9)
                run.font.name = "Consolas"
                code_lines = []
                i += 1
                continue
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("```"):
            flush_table()
            in_code = True
            code_lang = line.strip()[3:].strip()
            code_lines = []
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buffer.append(line)
            i += 1
            continue
        else:
            flush_table()

        if line.strip() == "---":
            add_formatted_paragraph(doc, "---")
            i += 1
            continue

        hl = heading_level(line)
        if hl:
            level, title = hl
            # Map # -> Title, ## -> Heading 1, etc.
            if level == 1:
                doc.add_heading(title, level=0)
            else:
                doc.add_heading(title, level=min(level - 1, 3))
            i += 1
            continue

        # Ordered list: "1. text" or duplicate "1." from MD quirks
        m_ol = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m_ol:
            add_formatted_paragraph(doc, f"{m_ol.group(1)}. {m_ol.group(2)}", style="List Number")
            i += 1
            continue

        if line.strip().startswith("- "):
            add_formatted_paragraph(doc, line.strip()[2:], style="List Bullet")
            i += 1
            continue

        if line.strip().startswith("  - "):
            add_formatted_paragraph(doc, line.strip()[4:], style="List Bullet 2")
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        add_formatted_paragraph(doc, line)
        i += 1

    flush_table()
    if in_code and code_lines:
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_lines))
        set_run_font(run, name="Consolas", size=9)

    doc.save(out_path)


def main():
    md_path = MD_PATH
    out_path = OUT_PATH
    if len(sys.argv) > 1:
        md_path = Path(sys.argv[1])
    if len(sys.argv) > 2:
        out_path = Path(sys.argv[2])
    text = md_path.read_text(encoding="utf-8")
    try:
        convert(text, out_path)
        print(f"Wrote {out_path}")
    except PermissionError:
        fallback = out_path.with_name(out_path.stem + "-updated.docx")
        convert(text, fallback)
        print(
            f"Could not write {out_path} (file open in Word?). Wrote {fallback} instead."
        )


if __name__ == "__main__":
    main()
