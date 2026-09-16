"""Generate C360 Customer 360 LWC DTO inventory Excel."""

import csv
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "c360LwcDtoInventory.xlsx"
OUTPUT_CSV = ROOT / "c360LwcDtoInventory.csv"

# OneDrive/SharePoint can inject [trash]/*.dat and customXml/* into .xlsx files,
# which causes Excel Online/desktop to hang on "Processing". Strip on write.
JUNK_PREFIXES = ("[trash]/", "customXml/")
JUNK_EXACT = {"docProps/custom.xml"}

HEADERS = [
    "DTO Object",
    "Field Name",
    "Type",
    "Requirement",
    "Page",
    "Component",
    "Section Title",
    "Data Source",
]

# UI panel / module headings (match LWC h2 and @api title defaults).
PORTFOLIO_SUMMARY_SECTION = {
    "totalAccounts": "Total accounts",
    "compositeHealthIndex": "Composite health index",
    "healthIndexDelta": "Composite health index",
    "signalsToAction": "Signals to action",
    "signalsChurnCount": "Signals to action",
    "signalsCrossSellCount": "Signals to action",
    "atRiskAccounts": "At-risk accounts",
    "revenueAtRisk": "Revenue at risk",
    "revenueAtRiskDelta": "Revenue at risk",
}

CROSS_SELL_SUMMARY_NET_BENEFIT = {
    "netBenefit",
    "grossUplift",
    "inScopeTransactions",
    "inScopeVolume",
    "declineCodeList",
    "cost",
    "tokenisationImpact",
}

CROSS_SELL_SUMMARY_ENV_DECLINES = {
    "merchantApprovalRate",
    "peerApprovalRate",
    "merchantVsPeer",
}

CROSS_SELL_SUMMARY_DECLINE_KPIS = {
    "valueOfDeclines",
    "countDeclines",
    "recoverableSplitMerchant",
    "recoverableSplitPeer",
}

ACCOUNT_DETAIL_CHURN_METRICS = {
    "churnScore",
    "churnScorePct",
    "txnTrend1m",
    "txnTrend3m",
    "txnTrend6m",
}

ACCOUNT_HIGHLIGHT_FIELDS = {
    "accountId",
    "merchantCount",
    "market",
}


def infer_section_title(dto: str, field: str, page: str, component: str) -> str:
    """Map a DTO field to the on-screen section heading where it is shown."""
    if dto == "PortfolioSummary" and field in PORTFOLIO_SUMMARY_SECTION:
        return PORTFOLIO_SUMMARY_SECTION[field]

    if dto == "CrossSellSummary":
        if field in CROSS_SELL_SUMMARY_NET_BENEFIT:
            return "Net benefit to merchant"
        if field in CROSS_SELL_SUMMARY_ENV_DECLINES:
            return "Account environment declines"
        if field in CROSS_SELL_SUMMARY_DECLINE_KPIS:
            return "Count & value of declines"

    if dto == "DeclineCode":
        return "Count & value of declines"

    if dto in ("CrossSellDetail",) and field == "schemeInterchangeBenefit":
        return "Net benefit to merchant"

    if dto == "CrossSellDetail" and field == "peerSegmentId":
        return "Account environment declines"

    if dto == "CrossSellDetail" and field == "likelihoodToAcquire":
        return "Net benefit to merchant"

    if dto == "AccountDetail" and field in ACCOUNT_DETAIL_CHURN_METRICS:
        return "Churn score & transaction trends"

    if dto == "AccountDetail" and field in ACCOUNT_HIGHLIGHT_FIELDS:
        return "Account highlight"

    if dto == "AccountDetail" and field == "crossSell":
        return "Revenue Boost cross-sell deep-dive"

    if dto == "AccountDetail" and field in (
        "positiveDrivers",
        "negativeDrivers",
        "driverDrilldown",
        "mids",
    ):
        return "Account-level drivers"

    if dto == "DriverRow":
        return "Account-level drivers"

    if dto == "DrilldownRow":
        return "Mid-level drivers"

    if dto == "MidRow":
        return "Account-level drivers"

    if dto in ("BASE_PRICE_PER_TXN",) or (
        dto == "CrossSellDetail" and field == "declineCodes"
    ):
        if dto == "BASE_PRICE_PER_TXN":
            return "Revenue Boost assumptions (sliders)"
        return "Count & value of declines"

    if dto == "DECLINE_CHART_COLORS":
        return "Count & value of declines"

    dto_defaults = {
        "UserContext": "Global header & sidebar",
        "PortfolioScope": "Enterprise Account Hub",
        "DataProvenance": "My accounts",
        "MOCK_ACCOUNT_IDS": "Navigation (account links)",
        "MaterialChangeBanner": "Material changes banner",
        "KPI_TILES": "KPI strip",
        "KpiTile": "KPI strip",
        "NEEDS_ACTION_ITEMS": "Needs action today",
        "NeedsActionItem": "Needs action today",
        "PORTFOLIO_HEALTH": "Portfolio health",
        "PortfolioHealth": "Portfolio health",
        "PipelineBar": "Portfolio health — cross-sell pipeline",
        "SIGNALS": "My signals",
        "Signal": "My signals",
        "ACCOUNTS": "My accounts",
        "Account": "My accounts",
        "ALERTS": "Active alerts",
        "Alert": "Active alerts",
        "ALERT_STATUS_OPTIONS": "Active alerts — status filters",
        "ALERT_INTEL": "Cross-selling intelligence",
        "CrossSellIntelligencePanel": "Cross-selling intelligence",
        "CrossSellIntelRow": "Cross-selling intelligence",
        "ALERT_DETAIL_BY_ID": "Alert detail",
        "AlertDetail": "Alert detail — header",
        "Driver": "Top drivers",
        "AlertCrossSellIntel": "Cross-sell intelligence",
        "AlertOutcome": "Record outcome",
        "CROSS_SELL_ROWS": "Open cross-sell opportunities",
        "CrossSellOpportunity": "Open cross-sell opportunities",
        "CHURN_ROWS": "Accounts by churn risk",
        "ChurnRiskRow": "Accounts by churn risk",
        "AccountSpotlight": "Account highlight",
        "AccountDetailInput": "Account 360",
        "ACCOUNT_DETAIL": "Account 360",
        "AccountDetail": "Account 360",
        "CrossSellDetail": "Revenue Boost cross-sell deep-dive",
        "CrossSellSummary": "Net benefit to merchant",
        "DashboardState": "App shell & navigation",
    }

    if dto == "AlertDetail" and field == "trajectory":
        return "6-month score trajectory"
    if dto == "AlertDetail" and field == "drivers":
        return "Top drivers"
    if dto in ("AlertDetail",) and field in ("score", "arr", "merchantCount", "account", "accountId", "id"):
        return "Alert detail — header"

    if dto in dto_defaults:
        return dto_defaults[dto]

    component_defaults = {
        "c360Dashboard": "Enterprise Account Hub",
        "c360MaterialBanner": "Material changes banner",
        "c360KpiStrip": "KPI strip",
        "c360KpiTile": "KPI strip",
        "c360NeedsAction": "Needs action today",
        "c360PortfolioHealth": "Portfolio health",
        "c360MySignals": "My signals",
        "c360SignalList": "My signals",
        "c360AccountsTable": "My accounts",
        "c360AlertCentre": "Alert Centre",
        "c360AlertDetail": "Alert detail",
        "c360ChurnTable": "Accounts by churn risk",
        "c360CrossSellTable": "Open cross-sell opportunities",
        "c360Account": "Account highlight",
        "c360AccountDetail": "Account 360",
        "c360PathwaysTable": "In-progress pathways",
        "c360ExportModal": "Export modal",
    }
    if component in component_defaults:
        return component_defaults[component]

    if "Deprecated" in page or "removed" in page.lower():
        return "Deprecated UI"
    return page


def with_section_title(row: tuple) -> tuple:
    dto, field, typ, requirement, page, component, data_source = row
    section = infer_section_title(dto, field, page, component)
    return (dto, field, typ, requirement, page, component, section, data_source)


def expand_rows(rows: list[tuple]) -> list[tuple]:
    return [with_section_title(row) for row in rows]

ROWS = [
    # Shared session context
    ("UserContext", "userId", "Id (string)", "Required", "Hub - Shared", "c360Dashboard", ""),
    ("UserContext", "userName", "string", "Required", "Hub - Shared", "c360Dashboard", ""),
    ("UserContext", "userInitials", "string", "Optional", "Hub - Shared", "c360Dashboard", ""),
    ("UserContext", "role", "string", "Required", "Hub - Shared", "c360Dashboard", ""),
    ("UserContext", "segment", "string", "Optional", "Hub - Shared", "c360Dashboard", ""),
    ("UserContext", "pilotGroupId", "string", "Optional", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("UserContext", "isPilotRm", "boolean", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("UserContext", "managerVisibility", "boolean", "Optional", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("PortfolioScope", "accountCount", "integer", "Required", "Hub - Overview", "c360Dashboard", ""),
    ("PortfolioScope", "asOfDate", "date", "Required", "Hub - Overview", "c360Dashboard", ""),
    ("PortfolioScope", "currency", "string", "Optional", "Hub - Overview", "c360Dashboard", ""),
    ("DataProvenance", "sourceChain", "string", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("DataProvenance", "refreshIntervalMinutes", "integer", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("DataProvenance", "lastRefreshedAt", "datetime", "Optional", "Hub - Overview", "c360AccountsTable", ""),
    ("MOCK_ACCOUNT_IDS", "(collection)", "object map", "Required", "Hub - Shared", "c360Dashboard", ""),
    ("MOCK_ACCOUNT_IDS", "accountName → salesforceId", "Id", "Required", "Hub - Shared", "c360Dashboard", ""),
    # Material banner
    ("MaterialChangeBanner", "message", "string", "Required", "Hub - Overview", "c360MaterialBanner", ""),
    ("MaterialChangeBanner", "accountsMovedToAtRisk", "integer", "Optional", "Hub - Overview", "c360MaterialBanner", ""),
    ("MaterialChangeBanner", "newChurnSignals", "integer", "Optional", "Hub - Overview", "c360MaterialBanner", ""),
    ("MaterialChangeBanner", "crossSellPipelineDelta", "Money", "Optional", "Hub - Overview", "c360MaterialBanner", ""),
    # KPI tiles
    ("KPI_TILES", "(collection)", "KpiTile[]", "Required", "Hub - Overview", "c360KpiStrip", ""),
    ("KpiTile", "label", "string", "Required", "Hub - Overview", "c360KpiTile", ""),
    ("KpiTile", "value", "string", "Required", "Hub - Overview", "c360KpiTile", ""),
    ("KpiTile", "detail", "string", "Optional", "Hub - Overview", "c360KpiTile", ""),
    ("KpiTile", "attention", "boolean", "Optional", "Hub - Overview", "c360KpiTile", ""),
    ("KpiTile", "variant", "enum (default|attention)", "Optional", "Hub - Overview", "c360KpiTile", ""),
    ("PortfolioSummary", "totalAccounts", "integer", "Required", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "compositeHealthIndex", "decimal", "Required", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "healthIndexDelta", "decimal", "Optional", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "signalsToAction", "integer", "Required", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "signalsChurnCount", "integer", "Optional", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "signalsCrossSellCount", "integer", "Optional", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "atRiskAccounts", "integer", "Required", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "revenueAtRisk", "Money|string", "Required", "Hub - Overview", "c360KpiStrip", ""),
    ("PortfolioSummary", "revenueAtRiskDelta", "decimal", "Optional", "Hub - Overview", "c360KpiStrip", ""),
    # Needs action
    ("NEEDS_ACTION_ITEMS", "(collection)", "NeedsActionItem[]", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "id", "string", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "severity", "enum (now|day|month)", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "severityLabel", "string", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "heading", "string", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "body", "string", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "meta", "string", "Optional", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "accountId", "Id", "Required (prod)", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "accountName", "string", "Required", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "buttonLabel", "string", "Optional", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "primary", "boolean", "Optional", "Hub - Overview", "c360NeedsAction", ""),
    ("NeedsActionItem", "sourceView", "enum (accounts|churn|crosssell)", "Optional", "Hub - Overview", "c360NeedsAction", ""),
    # Portfolio health
    ("PORTFOLIO_HEALTH", "(collection)", "PortfolioHealth", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PortfolioHealth", "healthy", "integer", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PortfolioHealth", "watch", "integer", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PortfolioHealth", "atRisk", "integer", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PortfolioHealth", "pipeline", "PipelineBar[]", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PipelineBar", "name", "string", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PipelineBar", "detail", "string", "Optional", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PipelineBar", "value", "Money|string", "Required", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PipelineBar", "width", "string", "Optional", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PipelineBar", "productCode", "string", "Optional", "Hub - Overview", "c360PortfolioHealth", ""),
    ("PipelineBar", "opportunityCount", "integer", "Optional", "Hub - Overview", "c360PortfolioHealth", ""),
    # Signals
    ("SIGNALS", "(collection)", "Signal[]", "Required", "Hub - Overview", "c360MySignals", ""),
    ("Signal", "id", "string", "Required", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "category", "enum (Cross-sell|Churn|FX)", "Required", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "badgeClass", "string", "Optional", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "title", "string", "Required", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "description", "string", "Required", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "signalStatus", "enum (open|actioned|dismissed)", "Required", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "accountId", "Id", "Optional", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "accountName", "string", "Optional", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "severity", "enum", "Optional", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "impactRevenue", "Money", "Optional", "Hub - Overview", "c360SignalList", ""),
    ("Signal", "triggeredAt", "datetime", "Optional", "Hub - Overview", "c360SignalList", ""),
    # Accounts table
    ("ACCOUNTS", "(collection)", "Account[]", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "id", "string", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "salesforceId", "Id", "Required (prod)", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "name", "string", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "industry", "string", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "volume", "string", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "healthIndex", "decimal", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "healthDelta", "string", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "health", "enum (Healthy|Watch|At risk)", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "healthClass", "string", "Optional", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "risk", "enum (Low|Medium|High)", "Required", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "riskClass", "string", "Optional", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "crossSell", "string", "Optional", "Hub - Overview", "c360AccountsTable", ""),
    ("Account", "lastActivityAt", "datetime", "Optional", "Hub - Overview", "c360AccountsTable", ""),
    # Alert centre
    ("ALERTS", "(collection)", "Alert[]", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "id", "string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "accountId", "Id", "Required (prod)", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "account", "string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "score", "decimal", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "alertDate", "date|string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "suppression", "string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "rmNotification", "string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "email", "string", "Optional", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "status", "enum", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "severity", "enum", "Optional", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("Alert", "revenueAtRisk", "Money|string", "Optional", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("ALERT_STATUS_OPTIONS", "(collection)", "string[]", "Optional", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("ALERT_INTEL", "(collection)", "object map", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("CrossSellIntelligencePanel", "pipelineValue", "Money", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("CrossSellIntelligencePanel", "averageConfidence", "decimal", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("CrossSellIntelligencePanel", "opportunities", "CrossSellIntelRow[]", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("CrossSellIntelRow", "accountName", "string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("CrossSellIntelRow", "suggestedProduct", "string", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    ("CrossSellIntelRow", "confidencePercent", "decimal", "Required", "Hub - Alert Centre", "c360AlertCentre", ""),
    # Alert detail
    ("ALERT_DETAIL_BY_ID", "(collection)", "object map", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "id", "string", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "accountId", "Id", "Required (prod)", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "account", "string", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "score", "decimal", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "arr", "Money|string", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "merchantCount", "integer", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "trajectory", "decimal[]", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertDetail", "drivers", "Driver[]", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("Driver", "label", "string", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("Driver", "impact", "string", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("Driver", "width", "string", "Optional", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertCrossSellIntel", "pipeline", "Money|string", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertCrossSellIntel", "confidence", "decimal", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertCrossSellIntel", "products", "string[]", "Required", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertOutcome", "interactionDate", "date", "Required (write)", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertOutcome", "interactionType", "enum", "Required (write)", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertOutcome", "outcome", "enum", "Required (write)", "Hub - Alert Detail", "c360AlertDetail", ""),
    ("AlertOutcome", "followUpActions", "string", "Optional (write)", "Hub - Alert Detail", "c360AlertDetail", ""),
    # Cross-sell table
    ("CROSS_SELL_ROWS", "(collection)", "CrossSellOpportunity[]", "Required", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "accountId", "Id", "Required (prod)", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "account", "string", "Required", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "product", "string", "Required", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "propensity", "enum (High|Medium|Low)", "Required", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "propensityClass", "string", "Optional", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "uplift", "Money|string", "Required", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    ("CrossSellOpportunity", "driver", "string", "Required", "Hub - Cross-Sell", "c360CrossSellTable", ""),
    # Churn table
    ("CHURN_ROWS", "(collection)", "ChurnRiskRow[]", "Required", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "accountId", "Id", "Required (prod)", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "account", "string", "Required", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "risk", "enum (High|Medium|Low)", "Required", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "riskClass", "string", "Optional", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "txnChange", "string", "Required", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "driver", "string", "Required", "Hub - Churn", "c360ChurnTable", ""),
    ("ChurnRiskRow", "churnScore", "decimal", "Optional", "Hub - Churn", "c360ChurnTable", ""),
    # Account record page wrapper
    ("AccountSpotlight", "recordId", "Id", "Required", "Account Record Page", "c360Account", ""),
    ("AccountSpotlight", "prefetchedData", "object", "Optional", "Account Record Page", "c360Account", ""),
    ("AccountSpotlight", "sourceView", "enum (accounts|churn|crosssell)", "Optional", "Account Record Page", "c360Account", ""),
    ("AccountDetailInput", "account", "Account", "Required", "Account Record Page", "c360AccountDetail", ""),
    ("AccountDetailInput", "sourceView", "enum (accounts|churn|crosssell)", "Optional", "Account Record Page", "c360AccountDetail", ""),
    # Account detail root
    ("ACCOUNT_DETAIL", "(collection)", "object map", "Required", "Account Record Page", "c360AccountDetail", ""),
    ("AccountDetail", "accountId", "string", "Required", "Account Record Page", "c360AccountDetail", ""),
    ("AccountDetail", "merchantCount", "integer", "Required", "Account Record Page", "c360AccountDetail", ""),
    ("AccountDetail", "market", "string", "Required", "Account Record Page", "c360AccountDetail", ""),
    ("AccountDetail", "churnScore", "number", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "churnScorePct", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "txnTrend1m", "string", "Optional (not rendered)", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "txnTrend3m", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "txnTrend6m", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "positiveDrivers", "DriverRow[]", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "negativeDrivers", "DriverRow[]", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "driverDrilldown", "object map", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "mids", "MidRow[]", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("AccountDetail", "crossSell", "CrossSellDetail", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DriverRow", "name", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("DriverRow", "value", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    # normalizeDriverDrilldown() reshapes the seed rows into {id, value} before render
    ("DrilldownRow", "id", "string (seed: humanCustomerId)", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("DrilldownRow", "value", "string (seed: driver metric)", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("MidRow", "id", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("MidRow", "name", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("MidRow", "negName", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("MidRow", "negValue", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("MidRow", "posName", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("MidRow", "posValue", "string", "Required", "Account Record Page - Churn tab", "c360AccountDetail", ""),
    ("CrossSellDetail", "likelihoodToAcquire", "enum", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellDetail", "peerSegmentId", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellDetail", "summary", "CrossSellSummary", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellDetail", "schemeInterchangeBenefit", "number", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellDetail", "declineCodes", "DeclineCode[]", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "netBenefit", "number", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "grossUplift", "number", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "inScopeTransactions", "integer", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "declineCodeList", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "merchantApprovalRate", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "peerApprovalRate", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "merchantVsPeer", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "recoverableSplitMerchant", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "recoverableSplitPeer", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "valueOfDeclines", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "countDeclines", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "cost", "number", "Optional", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "tokenisationImpact", "number", "Optional", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("CrossSellSummary", "inScopeVolume", "number", "Optional", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "code", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "name", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "volume", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "count", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "countNum", "integer", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "color", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "shareMerchant", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "sharePeer", "string", "Optional", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "curable", "enum (Y|N)", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "curePct", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "eligible", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DeclineCode", "recoveredRevenue", "string", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("BASE_PRICE_PER_TXN", "(constant)", "number", "Required", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    ("DECLINE_CHART_COLORS", "(constant)", "string[]", "Optional", "Account Record Page - Cross-sell tab", "c360AccountDetail", ""),
    # Dashboard orchestrator state
    ("DashboardState", "activeView", "enum", "Required", "Hub - Shared", "c360Dashboard", ""),
    ("DashboardState", "selectedAlertId", "string", "Required", "Hub - Alert Detail", "c360Dashboard", ""),
    ("DashboardState", "selectedAccount", "Account", "Required", "Hub - Account inline", "c360Dashboard", ""),
    ("DashboardState", "accountSourceView", "enum (accounts|churn|crosssell)", "Required", "Hub - Account inline", "c360Dashboard", ""),
    ("DashboardState", "accountSubTab", "enum (ch|cx)", "Required", "Hub - Account inline", "c360Dashboard", ""),
    ("DashboardState", "signals", "Signal[]", "Required", "Hub - Overview", "c360Dashboard", ""),
    ("DashboardState", "alertStatusFilter", "string", "Required", "Hub - Alert Centre", "c360Dashboard", ""),
]

DEPRECATED_ROWS = [
    ("Signal", "pathway", "string", "Deprecated", "Hub - Overview (v5 removed)", "c360SignalList", ""),
    ("Signal", "isOpen", "boolean", "Deprecated", "Hub - Overview (v5 removed)", "c360SignalList", ""),
    ("Alert", "threshold", "decimal", "Deprecated", "Hub - Alert Centre (v5 removed)", "c360AlertCentre", ""),
    ("Alert", "rmTask", "string", "Deprecated", "Hub - Alert Centre (v5 removed)", "c360AlertCentre", ""),
    ("ChurnRiskRow", "pathway", "string", "Deprecated", "Hub - Churn (v5 removed)", "c360ChurnTable", ""),
    ("AccountDetail", "authRate", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("CrossSellDetail", "slices", "SliceRow[]", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("CrossSellDetail", "productDrivers", "ProductDriver[]", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("PipelineBar", "name", "FraudSight", "Deprecated", "Hub - Overview (v5 removed)", "c360PortfolioHealth", ""),
    ("PATHWAYS", "(collection)", "Pathway[]", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "id", "string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "account", "string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "name", "string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "stage", "string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "owner", "string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "date", "date|string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "status", "enum", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("Pathway", "scope", "enum (me|team)", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
    ("ExportRequest", "accountId", "Id", "Deprecated", "Hub (removed)", "c360ExportModal", ""),
    ("ExportRequest", "artefacts", "string[]", "Deprecated", "Hub (removed)", "c360ExportModal", ""),
    ("ExportModal", "accountName", "string", "Deprecated", "Hub (removed)", "c360ExportModal", ""),
    # SliceRow orphaned with its parent CrossSellDetail.slices - no getter consumes it
    ("SliceRow", "credential", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "channel", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "scheme", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "market", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "volume", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "transactions", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "approval", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "recMerchant", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "peerApproval", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    ("SliceRow", "recPeer", "string", "Deprecated", "Account Record Page - Cross-sell tab (v5 removed)", "c360AccountDetail", ""),
    # Pre-normalisation seed shape - superseded by DrilldownRow.id / .value
    ("DrilldownRow", "humanCustomerId", "string", "Deprecated", "Account Record Page - Churn tab (pre-normalisation)", "c360AccountDetail", ""),
    ("DrilldownRow", "authRateAvg6m", "string", "Deprecated", "Account Record Page - Churn tab (pre-normalisation)", "c360AccountDetail", ""),
    # Present in c360AccountDetailData.js but no getter consumes them
    ("AccountDetail", "chargebacks", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("AccountDetail", "servicing", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("AccountDetail", "volumeTrend", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("AccountDetail", "volumeTrendNote", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("AccountDetail", "authTrend", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("AccountDetail", "cbTrend", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("AccountDetail", "servicingNote", "string", "Deprecated", "Account Record Page - Overview tab (v5 removed)", "c360AccountDetail", ""),
    ("Account", "pathways", "integer", "Deprecated", "Hub - Overview (v5 removed)", "c360AccountsTable", ""),
    ("Pathway", "statusClass", "string", "Deprecated", "Hub - Pathways (removed)", "c360PathwaysTable", ""),
]


def write_sheet(ws, rows, title_note=None):
    header_fill = PatternFill("solid", fgColor="1B5297")
    header_font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"
    widths = [22, 28, 24, 16, 28, 34, 32, 36]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if title_note:
        ws.cell(row=1, column=len(HEADERS) + 1, value=title_note)


def _is_junk_zip_entry(name: str) -> bool:
    if name in JUNK_EXACT:
        return True
    return any(name.startswith(prefix) for prefix in JUNK_PREFIXES)


def sanitize_xlsx(src: Path, dest: Path) -> int:
    """Rewrite xlsx zip without SharePoint/OneDrive junk entries."""
    removed = 0
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if _is_junk_zip_entry(info.filename):
                removed += 1
                continue
            zout.writestr(info, zin.read(info.filename))
    return removed


def write_csv(path: Path) -> None:
    active = expand_rows(ROWS)
    deprecated = expand_rows(DEPRECATED_ROWS)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(active)
        writer.writerow([])
        writer.writerow(["--- Deprecated ---"] + [""] * (len(HEADERS) - 1))
        writer.writerows(deprecated)


def main():
    wb = Workbook()
    active = wb.active
    active.title = "DTO Inventory"
    write_sheet(active, expand_rows(ROWS))

    deprecated = wb.create_sheet("Deprecated")
    write_sheet(deprecated, expand_rows(DEPRECATED_ROWS))

    readme = wb.create_sheet("README")
    readme["A1"] = "C360 Customer 360 - LWC DTO Inventory"
    readme["A3"] = "Generated from c360MockData.js, c360AccountDetailData.js, moduleSpec.md, and data-schema.md"
    readme["A5"] = (
        "Section Title matches on-screen panel headings (h2 / module titles). "
        "Fill Data Source with upstream system/table/API (e.g. Snowflake, Data 360, Salesforce Account, Apex method)."
    )
    readme["A7"] = "Active hub pages: Overview, Alert Centre, Alert Detail, Churn, Cross-Sell."
    readme["A8"] = "Account deep-dives: Account Record Page only (c360Account → c360AccountDetail)."
    readme["A10"] = f"Active DTO rows: {len(ROWS)} | Deprecated rows: {len(DEPRECATED_ROWS)}"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    try:
        wb.save(tmp_path)
        stripped = sanitize_xlsx(tmp_path, OUTPUT)
        write_csv(OUTPUT_CSV)
    finally:
        tmp_path.unlink(missing_ok=True)

    print(f"Wrote {len(ROWS)} active + {len(DEPRECATED_ROWS)} deprecated rows to {OUTPUT}")
    print(f"Wrote CSV fallback to {OUTPUT_CSV}")
    if stripped:
        print(f"Stripped {stripped} SharePoint/OneDrive junk entries from xlsx")


if __name__ == "__main__":
    main()
