"""Convert a markdown file to .docx (headings, tables, lists, code, inline bold/code)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_inline_runs(paragraph, text: str) -> None:
    """Parse **bold**, *italic*, and `code` into runs."""
    pattern = re.compile(
        r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[0-9]+\]|[^*`\[]+)"
    )
    for part in pattern.findall(text):
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def is_separator_row(line: str) -> bool:
    s = line.strip().strip("|")
    return bool(re.match(r"^[\s\-:|]+$", s))


def parse_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = row[j] if j < len(row) else ""
            p = table.rows[i].cells[j].paragraphs[0]
            add_inline_runs(p, cell_text)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                set_cell_shading(table.rows[i].cells[j], "E7E6E6")
    doc.add_paragraph()


def md_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    list_buffer: list[tuple[str, str]] = []  # (type, text) bullet or number

    def flush_list() -> None:
        nonlocal list_buffer
        for kind, text in list_buffer:
            p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else "List Number")
            add_inline_runs(p, text)
        list_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if in_code:
            if stripped.startswith("```"):
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                in_code = False
                code_lines = []
                code_lang = ""
            else:
                code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("```"):
            flush_list()
            code_lang = stripped[3:].strip()
            if code_lang == "mermaid":
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    i += 1
                p = doc.add_paragraph()
                run = p.add_run(
                    "[Flow diagram — recreate in Confluence or draw.io if needed.]\n"
                    "Set up → Build in Cursor (HTML, schema, LWC on sample data) → "
                    "Sandbox on sample data → Connect live data → Sustain."
                )
                run.italic = True
                if i < len(lines):
                    i += 1
                continue
            in_code = True
            code_lines = []
            i += 1
            continue

        if is_table_row(line):
            flush_list()
            table_rows: list[list[str]] = []
            while i < len(lines) and is_table_row(lines[i]):
                if not is_separator_row(lines[i]):
                    table_rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, table_rows)
            continue

        if stripped in ("---", "***", "___"):
            flush_list()
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue

        if not stripped:
            flush_list()
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)", stripped)
        number = re.match(r"^(\d+)\.\s+(.+)", stripped)
        if bullet:
            list_buffer.append(("bullet", bullet.group(1)))
            i += 1
            continue
        if number:
            list_buffer.append(("number", number.group(2)))
            i += 1
            continue

        flush_list()

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)
        i += 1

    flush_list()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    md = root / "docs" / "cursorDevelopmentScope.md"
    out = root / "docs" / "cursorDevelopmentScope.docx"
    if len(sys.argv) >= 2:
        md = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        out = Path(sys.argv[2])
    md_to_docx(md, out)
    print(out)


if __name__ == "__main__":
    main()
